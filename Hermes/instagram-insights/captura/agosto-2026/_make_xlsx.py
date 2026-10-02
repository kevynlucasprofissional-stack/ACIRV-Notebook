# -*- coding: utf-8 -*-
"""Gera o XLSX final normalizando os 3 formatos de JSON de _acirv_agosto."""
import json, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from _insights_builder import parse  # reuso do parser do builder

BASE = os.path.dirname(os.path.abspath(__file__))
MASTER = os.path.join(BASE, '..', '_acirv_agosto_progresso.json')
OUT = os.path.join(BASE, '..', 'instagram_acirvoficial_insights_agosto_2026.xlsx')

master = json.load(open(MASTER, encoding='utf-8'))
cands = master['candidatos'] if isinstance(master, dict) else master
def code_of(c):
    if isinstance(c, dict):
        return c.get('shortcode') or c.get('code')
    return c if isinstance(c, str) else None
order = [code_of(c) for c in cands]
meta_by_code = {code_of(c): c for c in cands if isinstance(c, dict)}
saved = {os.path.basename(f)[:-5]: f
         for f in glob.glob(os.path.join(BASE, '*.json'))
         if not os.path.basename(f).startswith('_')}

METRIC_ORDER = ['visualizacoes', 'visualizadores', 'contas_meta_alcancadas', 'alcance',
                'contas_alcancadas', 'interacoes', 'interacoes_posts', 'curtidas', 'comentarios',
                'compartilhamentos', 'salvamentos', 'contas_com_engajamento', 'na_pagina_inicial',
                'no_perfil', 'de_outra_pessoa', 'atividade_do_perfil', 'visitas_ao_perfil',
                'toques_em_links_externos', 'toques_no_endereco_comercial', 'novos_seguidores',
                'respostas', 'promovido']

allcodes = order + [c for c in sorted(saved) if c not in order]
cols = ['#', 'shortcode', 'data_publicacao', 'hora_publicacao', 'tipo_midia', 'autor',
        'colaborativo', 'coautores', 'promovido', 'visualizacoes'] + \
       [k for k in METRIC_ORDER if k not in ('visualizacoes', 'promovido')] + \
       ['obs_ou_nota', 'coletado_em', 'font[', 'url', 'legenda']

wb = openpyxl.Workbook()

# ---------- aba Posts ----------
ws = wb.active
ws.title = 'Posts'
hdr = ['#', 'shortcode', 'data', 'hora', 'tipo_midia', 'autor', 'colaborativo', 'coautores',
       'promovido'] + [k for k in METRIC_ORDER if k != 'promovido'] + \
      ['obs_ou_nota', 'coletado_em', 'fonte', 'url', 'legenda']
ws.append(hdr)
hf = Font(bold=True, color='FFFFFF')
fill = PatternFill('solid', fgColor='305496')
for c in range(1, len(hdr) + 1):
    cell = ws.cell(row=1, column=c)
    cell.font = hf; cell.fill = fill; cell.alignment = Alignment(vertical='top')

idx = 0
summary = {'viz': 0, 'promovidos': 0, 'sem_json': []}
for code in allcodes:
    if code not in saved:
        summary['sem_json'].append(code); continue
    d = json.load(open(saved[code], encoding='utf-8'))
    mm = d.get('melhores_metricas') or {}
    idx += 1
    row = [idx, code, d.get('data_publicacao'), d.get('hora_publicacao'), d.get('tipo_midia'),
           d.get('autor'), d.get('colaborativo'), ', '.join(d.get('coautores') or []),
           bool(mm.get('promovido'))]
    for k in METRIC_ORDER:
        if k == 'promovido':
            continue
        row.append(mm.get(k))
    obs = d.get('observacao') or d.get('nota') or ''
    row += [obs, d.get('coletado_em'), d.get('fonte'), d.get('url'), d.get('legenda')]
    ws.append(row)
    v = mm.get('visualizacoes')
    if isinstance(v, (int, float)):
        summary['viz'] += v
    if mm.get('promovido'):
        summary['promovidos'] += 1

# ---------- aba Metricas_brutas ----------
wb2 = wb.create_sheet('Metricas_brutas')
wb2.append(['shortcode', 'secao', 'metrica', 'valor'])
for c in range(1, 5):
    wb2.cell(row=1, column=c).font = hf
    wb2.cell(row=1, column=c).fill = fill
for code in allcodes:
    if code not in saved:
        continue
    d = json.load(open(saved[code], encoding='utf-8'))
    raw = d.get('insights_raw')
    if raw:
        for r in parse(raw.split('\n')):
            wb2.append([code, r['secao'], r['nome_original'], r['valor_original']])
    else:
        for s in (d.get('secoes') or []):
            sec = s.get('secao', '')
            for pair in (s.get('metricas') or []):
                if isinstance(pair, (list, tuple)) and len(pair) >= 2:
                    wb2.append([code, sec, pair[0], pair[1]])

# ---------- aba Controle ----------
wb3 = wb.create_sheet('Controle')
wb3.append(['#', 'shortcode', 'data_grid_master', 'tipo_esperado', 'json_existe',
            'tem_insights_raw', 'tem_secoes', 'tem_visualizacoes', 'promovido',
            'status_data', 'status_insights', 'coletado_em', 'url'])
for c in range(1, 14):
    wb3.cell(row=1, column=c).font = hf
    wb3.cell(row=1, column=c).fill = fill
ctrl_idx = 0
for code in allcodes:
    m = meta_by_code.get(code, {})
    exists = code in saved
    d = json.load(open(saved[code], encoding='utf-8')) if exists else {}
    mm = d.get('melhores_metricas') or {}
    ctrl_idx += 1
    wb3.append([ctrl_idx, code, m.get('data_grid'), m.get('tipo_esperado'), exists,
                bool(d.get('insights_raw')), bool(d.get('secoes')),
                isinstance(mm.get('visualizacoes'), (int, float)), bool(mm.get('promovido')),
                d.get('status_data') or ('PENDENTE' if not exists else ''),
                d.get('status_insights') or ('PENDENTE' if not exists else ''),
                d.get('coletado_em'), d.get('url') or m.get('url')])

# larguras
for sheet in (ws, wb2, wb3):
    for col in range(1, sheet.max_column + 1):
        L = get_column_letter(col)
        mx = max((len(str(sheet.cell(row=r, column=col).value or '')) for r in range(1, min(sheet.max_row, 200) + 1)), default=8)
        sheet.column_dimensions[L].width = min(max(10, mx + 2), 60)
ws.freeze_panes = 'C2'
wb2.freeze_panes = 'A2'
wb3.freeze_panes = 'A2'

wb.save(OUT)
print('SAVED:', OUT)
print('Posts linhas:', ws.max_row - 1, '| Metricas_brutas linhas:', wb2.max_row - 1,
      '| Controle linhas:', wb3.max_row - 1)
print('SOMA viz:', summary['viz'], '| promovidos:', summary['promovidos'],
      '| sem_json:', summary['sem_json'])
