# -*- coding: utf-8 -*-
"""Gera o XLSX final: abas Posts / Metricas_brutas / Controle."""
import json, os, sys, datetime
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _insights_builder import parse, NORM
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

BASE = r"C:\Users\Kevyn Lucas\Downloads\_acirv_agosto"
MASTER = r"C:\Users\Kevyn Lucas\Downloads\_acirv_agosto_progresso.json"
OUT = r"C:\Users\Kevyn Lucas\Downloads\instagram_acirvoficial_insights_agosto_2026.xlsx"

master = json.load(open(MASTER, encoding="utf-8"))
cands = master["candidatos"]

ORDER = ["shortcode","url","data_publicacao","hora_publicacao","tipo_midia","colaborativo","coautores",
         "promovido","autor","legenda",
         "visualizacoes","visualizadores","alcance","contas_alcancadas","contas_meta_alcancadas",
         "interacoes","interacoes_posts","curtidas","comentarios","compartilhamentos","salvamentos",
         "contas_com_engajamento","novos_seguidores","na_pagina_inicial","no_perfil","de_outra_pessoa",
         "atividade_do_perfil","visitas_ao_perfil","toques_em_links_externos","toques_no_endereco_comercial",
         "status_data","status_insights","metricas_ausentes","coletado_em","fonte"]

wb = Workbook()
bold = Font(bold=True)
head_fill = PatternFill("solid", fgColor="DDEBF7")

# ---------- aba Posts ----------
ws = wb.active; ws.title = "Posts"
ws.append(["n"] + ORDER)
for c in ws[1]:
    c.font = bold; c.fill = head_fill
    c.alignment = Alignment(vertical="center")

regs = []
for i, cd in enumerate(cands, 1):
    d = json.load(open(os.path.join(BASE, cd["shortcode"] + ".json"), encoding="utf-8"))
    mm = d.get("melhores_metricas") or {}
    row = {"shortcode": d.get("shortcode")}
    for k in ORDER:
        if k == "shortcode": continue
        if k in ("colaborativo","promovido"):
            row[k] = bool(d.get(k)) if k=="colaborativo" else bool(mm.get("promovido"))
        elif k == "coautores":
            row[k] = ", ".join(d.get("coautores") or [])
        elif k == "metricas_ausentes":
            row[k] = "; ".join(d.get("metricas_ausentes") or [])
        elif k in mm:
            row[k] = mm[k]
        else:
            row[k] = d.get(k)
    ws.append([i] + [row.get(k) for k in ORDER])
    regs.append((cd["shortcode"], d, mm))

widths = {"url":60,"legenda":70,"fonte":52,"shortcode":14,"metricas_ausentes":40,"coletado_em":24}
for j,k in enumerate(["n"]+ORDER, 1):
    ws.column_dimensions[get_column_letter(j)].width = widths.get(k, 13)
ws.freeze_panes = "B2"

# ---------- aba Metricas_brutas ----------
ws2 = wb.create_sheet("Metricas_brutas")
ws2.append(["shortcode","data_publicacao","secao","nome_original","valor_original","mapeado_para"])
for c in ws2[1]: c.font = bold; c.fill = head_fill
nrows = 0
for sc, d, mm in regs:
    raw = d.get("insights_raw") or ""
    if raw.strip():
        rows = parse(raw.split("\n"))
    else:
        rows = [{"secao": s.get("secao",""), "nome_original": nm, "valor_original": vl}
                for s in (d.get("secoes") or []) for nm, vl in s.get("metricas", [])]
    for r in rows:
        mapped = NORM.get(r["nome_original"], "")
        ws2.append([sc, d.get("data_publicacao"), r["secao"], r["nome_original"], r["valor_original"], mapped])
        nrows += 1
for j,w in enumerate([14,14,16,30,16,24],1):
    ws2.column_dimensions[get_column_letter(j)].width = w
ws2.freeze_panes = "A2"

# ---------- aba Controle ----------
ws3 = wb.create_sheet("Controle")
def put(k, v):
    ws3.append([k, v])
    ws3.cell(row=ws3.max_row, column=1).font = bold

agora = datetime.datetime.now().astimezone().isoformat(timespec="seconds")
prom = [sc for sc,_,mm in regs if mm.get("promovido")]
colab = [sc for sc,d,_ in regs if d.get("colaborativo")]
ausentes = [(sc, "; ".join(d.get("metricas_ausentes") or [])) for sc,d,_ in regs if d.get("metricas_ausentes")]
sem_dt = [sc for sc,d,_ in regs if d.get("status_data") != "ok"]
datas = sorted(d.get("data_publicacao") for _,d,_ in regs if d.get("data_publicacao"))

put("conta", "@acirvoficial")
put("periodo", "2026-08-01 a 2026-08-31 (America/Sao_Paulo, UTC-03:00)")
put("total_posts_lista_mestre", len(cands))
put("total_posts_no_arquivo", len(regs))
put("total_metricas_brutas", nrows)
put("data_min_publicacao", datas[0] if datas else "")
put("data_max_publicacao", datas[-1] if datas else "")
put("posts_promovidos", len(prom))
put("shortcodes_promovidos", ", ".join(prom))
put("posts_colaborativos", len(colab))
put("shortcodes_colaborativos", ", ".join(colab))
put("posts_sem_data_confirmada", len(sem_dt))
put("posts_com_metricas_ausentes", len(ausentes))
put("gerado_em", agora)
put("gerado_por", "Hermes Agent - coleta via pagina individual + 'Ver insights'")
put("fonte_primaria", "URL individual do post; botao 'Ver insights'; dados exibidos na tela de Insights")
put("arquivos_origem", BASE)
put("observacoes", ("Dados brutos preservados linha a linha na aba Metricas_brutas. "
                    "'promovido' = aviso real 'dados de anúncios' na tela de insights. "
                    "Metricas ausentes = rotulo presente sem valor numerico exibido."))
for j,w in enumerate([34,90],1): ws3.column_dimensions[get_column_letter(j)].width = w

wb.save(OUT)
print("OK ->", OUT)
print("abas:", wb.sheetnames, "| posts:", len(regs), "| linhas brutas:", nrows)
