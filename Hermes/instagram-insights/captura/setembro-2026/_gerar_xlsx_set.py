# -*- coding: utf-8 -*-
"""Gera o XLSX de setembro/2026 a partir de _set_linhas.json."""
import json, os, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
SAIDA = r"C:\Users\Kevyn Lucas\Downloads\instagram_acirvoficial_insights_setembro_2026.xlsx"

COLS = [
    ("media_id", 20), ("shortcode", 14), ("data_publicacao", 15), ("hora_publicacao", 9),
    ("tipo_midia", 10), ("vinculo", 13), ("parceiro", 22), ("colaborativo", 12),
    ("coautores", 22), ("promovido", 10), ("url", 58),
    ("visualizacoes", 13), ("visualizadores", 14), ("alcance", 10), ("contas_meta_alcancadas", 16),
    ("interacoes", 11), ("interacoes_posts", 15), ("interacoes_reels", 13),
    ("curtidas", 11), ("comentarios", 12), ("compartilhamentos", 15), ("salvamentos", 11),
    ("contas_com_engajamento", 16), ("novos_seguidores", 14),
    ("na_pagina_inicial", 14), ("no_perfil", 11), ("de_outra_pessoa", 14),
    ("atividade_do_perfil", 15), ("visitas_ao_perfil", 14),
    ("toques_em_links_externos", 17), ("toques_no_endereco_comercial", 19),
    ("pct_seguidores_alcance", 14), ("pct_nao_seguidores_alcance", 16),
    ("pct_seguidores_interacoes", 14), ("pct_nao_seguidores_interacoes", 16),
    ("taxa_interacao_pct", 13), ("status_data", 12), ("metricas_ausentes", 30),
    ("notas", 34), ("coletado_em", 20), ("arquivo_origem", 16),
]
INT_COLS = {"visualizacoes", "visualizadores", "alcance", "contas_meta_alcancadas", "interacoes",
            "interacoes_posts", "interacoes_reels", "curtidas", "comentarios", "compartilhamentos",
            "salvamentos", "contas_com_engajamento", "novos_seguidores", "na_pagina_inicial",
            "no_perfil", "de_outra_pessoa", "atividade_do_perfil", "visitas_ao_perfil",
            "toques_em_links_externos", "toques_no_endereco_comercial"}


def main():
    linhas = json.load(open(os.path.join(BASE, "_set_linhas.json"), encoding="utf-8"))
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E79")

    # ---------------------------------------------------------------- Posts
    ws = wb.active
    ws.title = "Posts"
    ws.append([c for c, _ in COLS])
    for i, (nome, largura) in enumerate(COLS, start=1):
        c = ws.cell(row=1, column=i)
        c.font = head_font
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = largura

    for l in linhas:
        m = l["metricas"]
        viz = m.get("visualizacoes")
        inter = m.get("interacoes")
        taxa = round(inter / viz * 100, 2) if isinstance(viz, int) and viz and isinstance(inter, int) else None
        linha = []
        for nome, _ in COLS:
            if nome in m:
                v = m[nome]
            elif nome == "coautores":
                v = ", ".join(l["coautores"])
            elif nome == "taxa_interacao_pct":
                v = taxa
            elif nome == "metricas_ausentes":
                v = ", ".join(l["metricas_ausentes"])
            elif nome == "notas":
                v = l["notas"]
            elif nome == "colaborativo":
                v = "sim" if l["colaborativo"] else "nao"
            elif nome == "promovido":
                v = "sim" if l["promovido"] else "nao"
            else:
                v = l.get(nome, "")
            linha.append(v)
        ws.append(linha)

    for r in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in r:
            nome = COLS[c.column - 1][0]
            if nome in INT_COLS and isinstance(c.value, int):
                c.number_format = "#,##0"
            elif nome == "taxa_interacao_pct" and isinstance(c.value, (int, float)):
                c.number_format = "0.00"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(COLS)), ws.max_row)

    # -------------------------------------------------------- Metricas_brutas
    ws2 = wb.create_sheet("Metricas_brutas")
    ws2.append(["media_id", "shortcode", "secao", "rotulo", "valor"])
    for i, (nome, largura) in enumerate([("media_id", 20), ("shortcode", 14), ("secao", 16),
                                         ("rotulo", 34), ("valor", 18)], start=1):
        c = ws2.cell(row=1, column=i)
        c.font = head_font
        c.fill = head_fill
        ws2.column_dimensions[get_column_letter(i)].width = largura
    n_brutas = 0
    for l in linhas:
        for b in l["metricas_brutas"]:
            if isinstance(b, dict):
                sec, rot, val = b.get("secao", ""), b.get("nome_original", b.get("rotulo", "")), b.get("valor_original", b.get("valor", ""))
            else:
                sec, rot, val = (list(b) + ["", "", ""])[:3]
            ws2.append([l["media_id"], l["shortcode"], sec, rot, val])
            n_brutas += 1
    ws2.freeze_panes = "A2"

    # ---------------------------------------------------------------- Resumo
    ws3 = wb.create_sheet("Resumo")
    def agrega(nome, sub):
        viz = sum(x["metricas"].get("visualizacoes", 0) or 0 for x in sub)
        inter = sum(x["metricas"].get("interacoes", 0) or 0 for x in sub)
        curt = sum(x["metricas"].get("curtidas", 0) or 0 for x in sub)
        com = sum(x["metricas"].get("comentarios", 0) or 0 for x in sub)
        comp = sum(x["metricas"].get("compartilhamentos", 0) or 0 for x in sub)
        salv = sum(x["metricas"].get("salvamentos", 0) or 0 for x in sub)
        n = len(sub)
        return [nome, n, viz, inter, curt, com, comp, salv,
                round(viz / n, 1) if n else 0, round(inter / n, 1) if n else 0]

    cab = ["grupo", "posts", "soma_visualizacoes", "soma_interacoes", "soma_curtidas",
           "soma_comentarios", "soma_compartilhamentos", "soma_salvamentos",
           "media_visualizacoes", "media_interacoes"]
    ws3.append(cab)
    for i in range(1, len(cab) + 1):
        c = ws3.cell(row=1, column=i)
        c.font = head_font
        c.fill = head_fill
        ws3.column_dimensions[get_column_letter(i)].width = max(14, len(cab[i - 1]) + 3)
    prop = [l for l in linhas if l["vinculo"] == "proprio"]
    colab = [l for l in linhas if l["vinculo"] == "colaboracao"]
    ws3.append(agrega("proprio", prop))
    ws3.append(agrega("colaboracao", colab))
    ws3.append(agrega("TOTAL", linhas))

    ws3.append([])
    ws3.append(["TOP 10 por visualizacoes"])
    ws3.cell(row=ws3.max_row, column=1).font = Font(bold=True)
    top_cab = ["media_id", "shortcode", "data_publicacao", "vinculo", "parceiro",
               "visualizacoes", "interacoes", "curtidas", "comentarios", "taxa_interacao_pct"]
    ws3.append(top_cab)
    for c in ws3[ws3.max_row]:
        c.font = Font(bold=True)
    for l in sorted(linhas, key=lambda x: -(x["metricas"].get("visualizacoes", 0) or 0))[:10]:
        m = l["metricas"]
        viz, inter = m.get("visualizacoes"), m.get("interacoes")
        taxa = round(inter / viz * 100, 2) if isinstance(viz, int) and viz and isinstance(inter, int) else None
        ws3.append([l["media_id"], l["shortcode"], l["data_publicacao"], l["vinculo"],
                    l["parceiro"] or "-", viz, inter, m.get("curtidas"), m.get("comentarios"), taxa])

    for r in ws3.iter_rows(min_row=2, max_row=ws3.max_row):
        for c in r:
            if isinstance(c.value, (int, float)) and c.column >= 3:
                c.number_format = "#,##0.0" if r[0].value in ("proprio", "colaboracao", "TOTAL") else "#,##0"

    # ---------------------------------------------------------------- Controle
    ws4 = wb.create_sheet("Controle")
    ws4.append(["chave", "valor"])
    ws4["A1"].font = head_font
    ws4["A1"].fill = head_fill
    ws4["B1"].font = head_font
    ws4["B1"].fill = head_fill
    ws4.column_dimensions["A"].width = 32
    ws4.column_dimensions["B"].width = 62

    viz_tot = sum(l["metricas"].get("visualizacoes", 0) or 0 for l in linhas)
    inter_tot = sum(l["metricas"].get("interacoes", 0) or 0 for l in linhas)
    datas = sorted(l["data_publicacao"] for l in linhas if l["data_publicacao"])
    sem_viz = [l["media_id"] for l in linhas if "visualizacoes" not in l["metricas"]]
    ctrl = [
        ("total_posts_no_arquivo", len(linhas)),
        ("posts_proprios", len(prop)),
        ("posts_colaboracao", len(colab)),
        ("postos_promovidos", sum(1 for l in linhas if l["promovido"])),
        ("periodo_de", datas[0] if datas else ""),
        ("periodo_ate", datas[-1] if datas else ""),
        ("total_metricas_brutas", n_brutas),
        ("soma_visualizacoes", viz_tot),
        ("soma_interacoes", inter_tot),
        ("posts_sem_visualizacoes", len(sem_viz)),
        ("media_ids_sem_visualizacoes", ", ".join(sem_viz) or "-"),
        ("contas_distintas_parceiro", len({l["parceiro"] for l in colab if l["parceiro"]})),
        ("parceiros", ", ".join(sorted({l["parceiro"] for l in colab if l["parceiro"]}))),
        ("fonte", "painel de insights do Instagram (conta @acirvoficial, user_id 6027695835)"),
        ("url_insights", "https://www.instagram.com/insights/media/<media_id>/"),
        ("capturas_cruas", "_raw_01.json + _raw_b02..b39.jsonl (ver coluna arquivo_origem)"),
        ("datas_origem", "inventario do grid: _grid_inventario.json (snowflake shortcode->media_id->data, desvio ~24s)"),
        ("gerado_em", datetime.datetime.now().astimezone().isoformat(timespec="seconds")),
        ("gerado_por", "pipeline _build_set.py + _gerar_xlsx_set.py"),
        ("AVISO_cabecalho", "no painel, o par 'Curtidas <n>' do topo repete a contagem de COMENTARIOS; os valores desta planilha vem do bloco rotulado (Interacoes com posts/reels)"),
        ("AVISO_facebook", "a subsecao Facebook do painel (Visualizacoes/Reacoes/Comentarios) foi descartada por sobrescrever os numeros do Instagram"),
        ("AVISO_alcance", "o painel de setembro/2026 nao expoe mais 'Alcance'/'Contas alcancadas'; ha 'Visualizadores' e, em alguns posts, 'Contas Meta alcancadas'"),
    ]
    for k, v in ctrl:
        ws4.append([k, v])

    wb.save(SAIDA)
    print("gravado:", SAIDA)
    print("Posts:", len(linhas), "| Metricas_brutas:", n_brutas)
    print("proprios:", len(prop), "| colaboracao:", len(colab))
    print("soma visualizacoes:", viz_tot, "| soma interacoes:", inter_tot)


if __name__ == "__main__":
    main()
