"""Testes da contagem de publicacoes do feed da ACIRV (agosto e setembro/2026).

Trava o numero e as invariantes que impedem o erro de voltar. Roda sem instalar nada:

    python test_publicacoes_feed.py

Tambem funciona sob pytest, se estiver instalado (pytest test_publicacoes_feed.py).

O que estes testes protegem:
 - a contagem e um CONJUNTO, nao um total (mesma contagem com posts trocados reprova);
 - nada pode aparecer numa fonte sem estar na varredura do feed;
 - a data vale pela decodificacao do shortcode, nao pelo rotulo da fonte;
 - um defeito conhecido fica travado como teste: quando for corrigido, o teste
   falha de proposito e obriga a atualizar a expectativa.
"""
import csv
import datetime
import json
import os
import sys

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
import publicacoes_feed as pf  # noqa: E402

FONTES = pf.registro_fontes()

def _grid():
    return pf.fonte_grid(FONTES["grid"][1])

def _src(nome):
    leitor, caminho = FONTES[nome][0], FONTES[nome][1]
    return leitor(caminho)

def _do_mes(itens, mes):
    return {sc for sc in itens if pf.mes_de(sc) == mes}

# ------------------------------------------------------------- a referencia
def test_01_grid_carrega_com_204_itens():
    g = _grid()
    assert len(g) == 204, "a varredura do feed tem %d itens, esperado 204" % len(g)

def test_02_shortcode_e_media_id_sao_coerentes():
    """O id que veio na captura tem que reproduzir o shortcode. Pega id trocado/copiado errado."""
    itens = json.load(open(FONTES["grid"][1], encoding="utf-8"))["itens"]
    ruins = [x["sc"] for x in itens if x.get("media_id") and pf.mid2sc(x["media_id"])[:11] != x["sc"][:11]]
    assert not ruins, "media_id nao corresponde ao shortcode em: %s" % ruins

def test_03_data_decodificada_confere_com_a_data_do_grid():
    """Se o decodificador errasse, todo o resto do teste estaria medindo a coisa errada."""
    itens = json.load(open(FONTES["grid"][1], encoding="utf-8"))["itens"]
    ruins = []
    for x in itens:
        esperada = str(x["publicado_brt"])[:16]
        obtida = pf.sc_para_data(x["sc"]).strftime("%Y-%m-%d %H:%M")
        if esperada != obtida:
            ruins.append((x["sc"], esperada, obtida))
    assert not ruins, "data decodificada != data da varredura em: %s" % ruins[:5]

# ---------------------------------------------------------------- contagens
def test_04_setembro_tem_exatamente_116_posts_no_feed():
    n = len(_do_mes(_grid(), "2026-09"))
    assert n == 116, "setembro tem %d posts na varredura do feed, esperado 116" % n

def test_05_agosto_tem_exatamente_62_posts_no_feed():
    n = len(_do_mes(_grid(), "2026-08"))
    assert n == 62, "agosto tem %d posts na varredura do feed, esperado 62" % n

def test_06_borda_de_mes_em_brt():
    """01/10 pertence a outubro e nao pode vazar para setembro."""
    g = _do_mes(_grid(), "2026-10")
    assert len(g) == 1, "outubro tem %d posts na varredura, esperado 1" % len(g)
    assert len(_do_mes(_grid(), "2026-09")) == 116, "o post de 01/10 vazou para setembro"

def test_07_cada_post_pertence_a_um_unico_mes():
    itens = json.load(open(FONTES["grid"][1], encoding="utf-8"))["itens"]
    meses = {}
    for x in itens:
        meses.setdefault(pf.mes_de(x["sc"]), []).append(x["sc"])
    total = sum(len(v) for v in meses.values())
    assert total == len(itens), "post contado em dois meses"
    assert set(meses) == {"2026-06", "2026-07", "2026-08", "2026-09", "2026-10"}, sorted(meses)

# ------------------------------------------------------------ base canonica
def test_08_base_de_setembro_e_o_mesmo_conjunto_do_grid():
    """Nao basta 116 = 116: tem que ser os MESMOS posts."""
    base = set(_src("base_set"))
    grid = _do_mes(_grid(), "2026-09")
    so_base, so_grid = base - grid, grid - base
    assert not so_base, "a base tem post que o feed nao tem: %s" % sorted(so_base)
    assert not so_grid, "o feed tem post que a base nao tem: %s" % sorted(so_grid)
    assert len(base) == len(grid) == 116

def test_09_base_de_setembro_nao_tem_shortcode_repetido():
    scs = [r["shortcode"][:11] for r in csv.DictReader(open(FONTES["base_set"][1], encoding="utf-8"), delimiter=";")]
    assert len(scs) == len(set(scs)) == 116

def test_10_xlsx_da_base_confere_com_o_csv():
    import openpyxl
    x = FONTES["base_set"][1].replace(".csv", ".xlsx")
    ws = openpyxl.load_workbook(x, read_only=True).active
    assert ws.max_row == 117, "xlsx tem %d linhas (1 cabecalho + 116), esperado 117" % ws.max_row

# --------------------------------------------------- uniao: nada de fora
def test_11_nenhuma_fonte_traz_post_que_o_feed_nao_tem():
    """O teste que pega o caso 'o sweep perdeu um post'."""
    for periodo in ("2026-09", "2026-08"):
        r = pf.auditar(periodo)
        for nome, f in r["fontes"].items():
            if "erro" in f:
                continue
            assert not f["fora_da_referencia"], "%s (periodo %s) traz post fora do feed: %s" % (
                nome, periodo, f["fora_da_referencia"])

def test_12_varreduras_antigas_sao_subconjunto_do_sweep_final():
    full = set(_grid())
    for nome in ("grid_v1", "grid_v2"):
        antigos = set(_src(nome))
        assert not (antigos - full), "%s tem post que o sweep final perdeu: %s" % (nome, sorted(antigos - full))

def test_13_brutos_e_export_de_agosto_cabem_no_grid():
    g = set(_grid())
    for nome in ("brutos_ago", "export_ago"):
        fora = set(_src(nome)) - g
        assert not fora, "%s tem post fora do feed: %s" % (nome, sorted(fora))

def test_14_export_oficial_de_agosto_rotula_4_posts_com_data_errada():
    """Export ja trouxe 4 posts de 31/07 como se fossem de agosto. A data do shortcode manda."""
    esperados = {"Dbd6DYGA0y3", "DbdvooJOi2n", "DbdvwQLu4_o", "DbdwBOBOVJt"}
    exp = _src("export_ago")
    julho = {sc for sc in exp if pf.mes_de(sc) == "2026-07"}
    assert julho == esperados, "posts de julho no export de agosto mudaram: %s" % sorted(julho)
    agosto = {sc for sc in exp if pf.mes_de(sc) == "2026-08"}
    assert len(agosto) == 48
    assert not (agosto - set(_grid())), "export de agosto tem post que o feed nao tem"

# ------------------------------------------------- base canonica de agosto
def test_15_base_de_agosto_esta_completa_62_posts():
    """A divida dos 14 posts de agosto foi paga: a base canonica cobre os 62."""
    base = _src("base_ago26")
    g = _do_mes(_grid(), "2026-08")
    assert len(base) == 62, "a base de agosto tem %d linhas, esperado 62" % len(base)
    so_base = set(base) - g
    so_grid = g - set(base)
    assert not so_base and not so_grid, "base x feed: so_base=%s so_grid=%s" % (sorted(so_base), sorted(so_grid))

def test_16_base_de_agosto_tem_4_datas_publicadas_erradas():
    base = _src("base_08_09")
    errados = sorted(sc for sc in base if pf.mes_de(sc) == "2026-07" and str(base[sc].get("rotulo", ""))[:7] == "2026-08")
    assert errados == ["Dbd6DYGA0y3", "DbdvooJOi2n", "DbdvwQLu4_o", "DbdwBOBOVJt"], errados

# ------------------------------------------------------------------- buracos
def test_17_keniasleite_entra_na_contagem_mas_sem_views():
    """Reel colaborativo: lacuna do painel. Ausencia de metrica nunca vira zero."""
    base = _src("base_set")
    linha = None
    for r in csv.DictReader(open(FONTES["base_set"][1], encoding="utf-8"), delimiter=";"):
        if r["shortcode"][:11] == "DdFXhdwp0L9":
            linha = r
    assert linha is not None, "o reel colaborativo sumiu da base de setembro"
    assert linha.get("visualizacoes", "").strip() == "", "views do keniasleite veio preenchida: %r" % linha.get("visualizacoes")
    assert linha.get("status_validacao") == "colab_metricas_principais_indisponiveis_no_painel"

# ------------------------------------------------------- o auditor morde
def test_18_bateria_negativa_do_auditor_passa():
    ok, falhas = pf.autoteste()
    assert not falhas, "a bateria negativa falhou: %s" % falhas
    assert ok == 8, "esperava 8 casos na bateria negativa, rodou %d" % ok

def test_19_auditoria_reprova_quando_injetam_defeito():
    import tempfile
    tmp = tempfile.mkdtemp(prefix="teste_audit_")
    scs = sorted(_do_mes(_grid(), "2026-09"))[:3]
    gp = os.path.join(tmp, "grid.json")
    json.dump({"itens": [{"sc": s, "publicado_brt": pf.sc_para_data(s).strftime("%Y-%m-%d %H:%M:%S"),
                          "dono_grid": "acirvoficial", "tipo_grid": "p"} for s in scs]}, open(gp, "w", encoding="utf-8"))
    bp = os.path.join(tmp, "base.csv")
    with open(bp, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["shortcode", "data_publicacao", "vinculo"])
        for s in scs[:2]:
            w.writerow([s, pf.sc_para_data(s).strftime("%Y-%m-%d"), "proprio"])
    r = pf.auditar("2026-09", fontes={
        "grid": (pf.fonte_grid, gp, "grid de teste", "referencia", set()),
        "base": (pf.fonte_csv, bp, "base de teste", "base", {"2026-09"})})
    assert r["veredito"] == "REPROVADO", "o auditor aprovou uma base com 1 post faltando"

def test_20_veredito_limpo_quando_nada_falta():
    r = pf.auditar("2026-09")
    assert not r["divergencias"], "a auditoria de setembro acusou: %s" % r["divergencias"]
    assert r["veredito"] == "COMPLETO"

def test_21_base_de_agosto_tem_as_14_linhas_capturadas_no_painel():
    """As 14 publicacoes que faltavam entraram com metrica do painel post a post."""
    novas = {"DcbD8kGAZY3", "DcbECTUAleR", "DcbEHlnA9uZ", "DcbEPp2gwNj", "DcbmO6GPfhb",
             "DcdwpbfPsgs", "DchK9ycv_Tu", "Dchm6kFDjxS", "DclV0zcAUvH", "DclVRnSg4Sm",
             "DclVejgguEv", "DclVoHkAc3R", "DcmMX4pvjg7", "Dcrgb9wP9LW"}
    base = _src("base_ago26")
    assert not (novas - set(base)), "as 14 capturadas nao entraram: %s" % sorted(novas - set(base))
    linhas = list(csv.DictReader(open(FONTES["base_ago26"][1], encoding="utf-8"), delimiter=";"))
    aus = {r["shortcode"]: r["metricas_ausentes"] for r in linhas if r["shortcode"] in novas}
    sem = [k for k in ("DclV0zcAUvH", "DclVRnSg4Sm", "DclVejgguEv", "DclVoHkAc3R") if aus.get(k, "x") != ""]
    assert not sem, "linha de post proprio com ausente indevido: %s" % sem
    assert aus["DcbECTUAleR"] == "visualizadores", aus["DcbECTUAleR"]
    assert aus["Dcrgb9wP9LW"] == "atividade_do_perfil;novos_seguidores", aus["Dcrgb9wP9LW"]
    pro = {r["shortcode"]: r["promovido"] for r in linhas if r["shortcode"] in novas}
    assert pro["DcbECTUAleR"] == "sim" and pro["DclV0zcAUvH"] == "nao_informado", pro


def test_22_xlsx_e_csv_da_base_de_agosto_contam_a_mesma_coisa():
    import openpyxl
    caminho = FONTES["base_ago26"][1]
    ws = openpyxl.load_workbook(caminho.replace(".csv", ".xlsx")).active
    linhas = sum(1 for _ in csv.DictReader(open(caminho, encoding="utf-8"), delimiter=";"))
    assert linhas == 62, linhas
    assert ws.max_row == linhas + 1, "xlsx tem %d linhas, esperado %d" % (ws.max_row, linhas + 1)
    assert ws.max_column == 34


def test_23_split_do_legado_julho_extrato_e_setembro_subconjunto():
    """138 linhas = 52 de agosto (48 reais + 4 de 31/07) + 86 de setembro."""
    leg = list(csv.DictReader(open(FONTES["base_08_09"][1], encoding="utf-8"), delimiter=";"))
    assert len(leg) == 138, len(leg)
    julho = sorted(r["shortcode"] for r in leg if pf.mes_de(r["shortcode"]) == "2026-07")
    assert julho == ["Dbd6DYGA0y3", "DbdvooJOi2n", "DbdvwQLu4_o", "DbdwBOBOVJt"], julho
    setembro = {r["shortcode"] for r in leg if pf.mes_de(r["shortcode"]) == "2026-09"}
    assert len(setembro) == 86 and setembro <= set(_src("base_set")), "setembro do legado nao e subconjunto da canonica"
    cofre = os.path.dirname(os.path.dirname(FONTES["base_08_09"][1]))
    extrato = os.path.join(cofre, "Hermes", "instagram-insights", "auditorias", "julho-2026-31-do-export-master.csv")
    linhas = list(csv.DictReader(open(extrato, encoding="utf-8"), delimiter=";"))
    assert len(linhas) == 4, len(linhas)
    assert {r["data_publicacao"] for r in linhas} == {"2026-07-31"}
    assert {r["periodo_fonte"] for r in linhas} == {"2026-07"}


def test_24_legado_nunca_reprova_e_o_dia_errado_aparece():
    """Classe legado so gera aviso; o dia errado de agosto fica visivel no relatorio."""
    r = pf.auditar("2026-08")
    assert r["veredito"] == "COMPLETO", r["divergencias"]
    assert not r["divergencias"], r["divergencias"]
    assert r["fontes"]["base_08_09"]["classe"] == "legado", r["fontes"]["base_08_09"]["classe"]
    assert r["fontes"]["base_08_09"]["dia_errado"] == ["Db-7_-OOxRu"], r["fontes"]["base_08_09"]["dia_errado"]
    assert r["fontes"]["base_ago26"]["dia_errado"] == [] and r["fontes"]["base_ago26"]["total"] == 62


def test_25_base_de_legendas_cobre_exatamente_os_116_posts():
    """Regra do LEIA-ME: legenda e campo obrigatorio de toda postagem."""
    leg = _src("legendas_set")
    base = _src("base_set")
    assert set(leg) == set(base), "conjunto de shortcodes da legenda difere da base canonica"
    assert len(leg) == 116, len(leg)
    vazias = [sc for sc, d in leg.items() if not (d.get("legenda") or "").strip()]
    assert not vazias, "posts sem legenda: %s" % vazias
    curtas = [sc for sc, d in leg.items() if len(d["legenda"]) < 15]
    assert not curtas, "legenda curta demais para ser real: %s" % curtas


def test_26_o_portao_reprova_legenda_em_branco():
    import tempfile
    tmp = tempfile.mkdtemp(prefix="teste_legenda_")
    lp = os.path.join(tmp, "leg.csv")
    leg = _src("legendas_set")
    with open(lp, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["shortcode", "data_publicacao", "hora_publicacao", "vinculo", "tipo_midia", "legenda"])
        for i, (sc, d) in enumerate(sorted(leg.items())):
            w.writerow([sc, d["rotulo"].split(" ")[0], d["rotulo"].split(" ")[1], d["dono"], d["tipo"],
                        "" if i == 5 else d["legenda"]])
    r = pf.auditar("2026-09", fontes={
        "grid": (pf.fonte_grid, FONTES["grid"][1], "grid real", "referencia", set()),
        "leg": (pf.fonte_legendas, lp, "legenda de teste", "legenda", {"2026-09"})})
    assert r["veredito"] == "REPROVADO", "o auditor aprovou uma base com legenda em branco"
    assert r["fontes"]["leg"]["sem_legenda"], r["fontes"]["leg"]


def test_27_o_portao_reprova_post_faltando_na_base_de_legendas():
    import tempfile
    tmp = tempfile.mkdtemp(prefix="teste_legenda_")
    lp = os.path.join(tmp, "leg.csv")
    leg = _src("legendas_set")
    with open(lp, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["shortcode", "data_publicacao", "hora_publicacao", "vinculo", "tipo_midia", "legenda"])
        for sc, d in sorted(leg.items())[:-1]:
            w.writerow([sc, d["rotulo"].split(" ")[0], d["rotulo"].split(" ")[1], d["dono"], d["tipo"], d["legenda"]])
    r = pf.auditar("2026-09", fontes={
        "grid": (pf.fonte_grid, FONTES["grid"][1], "grid real", "referencia", set()),
        "leg": (pf.fonte_legendas, lp, "legenda de teste", "legenda", {"2026-09"})})
    assert r["veredito"] == "REPROVADO", "o auditor aprovou uma base de legendas com 1 post faltando"


def test_28_csv_e_xlsx_da_base_de_legendas_contam_a_mesma_coisa():
    import openpyxl
    csv_p = FONTES["legendas_set"][1]
    xlsx_p = csv_p[:-4] + ".xlsx"
    linhas = list(csv.DictReader(open(csv_p, encoding="utf-8"), delimiter=";"))
    ws = openpyxl.load_workbook(xlsx_p, read_only=True).active
    vals = list(ws.iter_rows(values_only=True))
    assert len(vals) - 1 == len(linhas) == 116, (len(vals), len(linhas))
    assert list(vals[0]) == list(linhas[0].keys()), "cabecalho do xlsx difere do csv"
    assert {r["shortcode"] for r in linhas} == {v[3] for v in vals[1:]}


# ------------------------------------------------------------------- runner
def test_29_toda_linha_das_bases_tem_media_id_que_fecha_com_o_shortcode():
    """Coluna de juncao nao pode sair em branco: a primeira montagem de agosto deixou
    48 das 62 linhas sem media_id e referencia cruzada por media_id quebrava calada."""
    import csv as _csv
    bases = os.path.normpath(os.path.join(AQUI, "..", "..", "..", "85-Bases-e-Consultas"))
    for mes in ("2026-08", "2026-09"):
        caminho = os.path.join(bases, "Instagram-Publicacoes-%s.csv" % mes)
        linhas = list(_csv.DictReader(open(caminho, encoding="utf-8"), delimiter=";"))
        assert linhas, mes
        vazias = [r["shortcode"] for r in linhas if not (r.get("media_id") or "").strip()]
        assert vazias == [], "%s: %d linha(s) sem media_id: %s" % (mes, len(vazias), vazias[:5])
        ruins = [r["shortcode"] for r in linhas
                 if pf.mid2sc(str(r["media_id"])) != r["shortcode"]]
        assert ruins == [], "%s: media_id nao fecha com o shortcode: %s" % (mes, ruins[:5])


def test_30_base_de_legendas_junta_1_para_1_com_a_base_de_publicacoes():
    """A legenda serve para dizer do que trata o post: se o conjunto de media_id nao for
    identico ao da base de publicacoes, existe post sem legenda ou legenda orfa."""
    import csv as _csv
    bases = os.path.normpath(os.path.join(AQUI, "..", "..", "..", "85-Bases-e-Consultas"))
    pubs = list(_csv.DictReader(
        open(os.path.join(bases, "Instagram-Publicacoes-2026-09.csv"), encoding="utf-8"), delimiter=";"))
    legs = list(_csv.DictReader(
        open(os.path.join(bases, "Instagram-Legendas-2026-09.csv"), encoding="utf-8"), delimiter=";"))
    mp = set(r["media_id"] for r in pubs)
    ml = set(r["media_id"] for r in legs)
    assert mp == ml, "sem legenda: %s | orfas: %s" % (sorted(mp - ml)[:5], sorted(ml - mp)[:5])
    assert len(legs) == 116
    curtas = [r["media_id"] for r in legs if len((r.get("legenda") or "").strip()) < 10]
    assert curtas == [], "legenda suspeita (curta demais): %s" % curtas[:5]


def test_31_base_de_legendas_de_agosto_fecha_com_a_de_publicacoes():
    """Agosto, igual a setembro: mesma quantidade, mesmo conjunto, e so o post
    declarado como sem legenda por natureza pode estar vazio."""
    import csv as _csv
    import os as _os
    bases = _os.path.normpath(_os.path.join(AQUI, "..", "..", "..", "85-Bases-e-Consultas"))
    pub = list(_csv.DictReader(open(_os.path.join(bases, "Instagram-Publicacoes-2026-08.csv"),
                                    encoding="utf-8-sig"), delimiter=";"))
    leg = list(_csv.DictReader(open(_os.path.join(bases, "Instagram-Legendas-2026-08.csv"),
                                    encoding="utf-8-sig"), delimiter=";"))
    assert len(pub) == 62, "base de agosto deveria ter 62 linhas, tem %d" % len(pub)
    assert len(leg) == 62, "base de legendas de agosto deveria ter 62 linhas, tem %d" % len(leg)
    assert set(r["shortcode"] for r in pub) == set(r["shortcode"] for r in leg), "conjuntos de shortcode divergem"
    assert set(r["media_id"] for r in pub) == set(r["media_id"] for r in leg), "conjuntos de media_id divergem"
    assert len(set(r["media_id"] for r in leg)) == 62, "media_id repetido na base de legendas"
    vazias = [r["shortcode"] for r in leg if not (r["legenda"] or "").strip()]
    assert set(vazias) <= pf.SEM_LEGENDA_POR_NATUREZA, \
        "legenda vazia fora da lista de posts sem legenda por natureza: %s" % vazias
    assert all(int(r["legenda_chars"]) == len(r["legenda"]) for r in leg), "legenda_chars nao fecha com o texto"

def test_32_todo_mes_com_base_de_publicacoes_tem_base_de_legendas():
    """O 'por padrao' virado invariante: nao existe mes com base canonica de
    publicacoes e sem base canonica de legendas. Se um mes novo entrar sem
    legenda, este teste reprova -- e o padrao deixa de depender de memoria."""
    import os as _os
    import re as _re
    bases = _os.path.normpath(_os.path.join(AQUI, "..", "..", "..", "85-Bases-e-Consultas"))
    pub = sorted(a for a in _os.listdir(bases)
                 if _re.fullmatch(r"Instagram-Publicacoes-\d{4}-\d{2}\.csv", a))
    assert pub, "nenhuma base canonica de publicacoes encontrada"
    faltando = []
    for a in pub:
        mes = _re.search(r"(\d{4}-\d{2})", a).group(1)
        if not _os.path.exists(_os.path.join(bases, "Instagram-Legendas-%s.csv" % mes)):
            faltando.append(mes)
    assert not faltando, "mes(es) com base de publicacoes e sem base de legendas: %s" % faltando
    # e o registro de fontes precisa exigir legenda em cada um desses meses
    fontes = pf.registro_fontes()
    exigidos = set()
    for esp in fontes.values():
        if esp[3] == "legenda":
            exigidos |= set(esp[4])
    nao_exigidos = [m for m in (_re.search(r"(\d{4}-\d{2})", a).group(1) for a in pub) if m not in exigidos]
    assert not nao_exigidos, "mes(es) com base de legendas que o portao nao exige: %s" % nao_exigidos

def main():
    testes = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    ok, falhas = 0, []
    for nome, f in testes:
        try:
            f()
            print("  PASSOU  %s" % nome)
            ok += 1
        except AssertionError as e:
            print("  FALHOU  %s\n          %s" % (nome, e))
            falhas.append(nome)
        except Exception as e:
            print("  ERRO    %s\n          %r" % (nome, e))
            falhas.append(nome)
    print("\n%d/%d testes passaram" % (ok, len(testes)))
    return 0 if not falhas else 1

if __name__ == "__main__":
    sys.exit(main())
