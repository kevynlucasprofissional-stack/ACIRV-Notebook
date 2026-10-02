"""Fonte unica da verdade para CONTAGEM de publicacoes do feed do Instagram da ACIRV.

Regra central (aprendida na marra): a contagem de um periodo e o CONJUNTO de
shortcodes da varredura do grid do perfil, nunca o total de um export.
Toda outra fonte e conferida como SUBCONJUNTO dessa referencia; post de fora
vira DIVERGENCIA e reprova a auditoria.

Tres regras que nasceram de erro real:
 1. Nunca comparar por media_id entre fontes: o export oficial vem sem media_id
    ou em outro espaco de id. A chave universal e o shortcode (11 chars).
 2. Nunca confiar na data rotulada: a data verdadeira sai da decodificacao do
    shortcode (validada contra as 204 datas do grid). Export ja apareceu com
    4 posts de 31/07 rotulados como agosto.
 3. Base nao fecha por contagem, fecha por conjunto: mesma contagem com posts
    trocados continua errada.

Uso:
    python publicacoes_feed.py auditar [--periodo 2026-09] [--periodo 2026-08]
    python publicacoes_feed.py autoteste
"""
import csv
import datetime
import json
import os
import sys

AQUI = os.path.dirname(os.path.abspath(__file__))
COFRE = os.path.abspath(os.environ.get("COFRE_INSTAGRAM") or os.path.join(AQUI, "..", "..", ".."))
ALFA = "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789-_"
EPOCH_MS = 1314220021721
OFFSET_S = -10800  # BRT
MESES = {"2026-06": "junho", "2026-07": "julho", "2026-08": "agosto", "2026-09": "setembro", "2026-10": "outubro"}

# ---------------------------------------------------------------- id <-> data
def sc2mid(sc):
    n = 0
    for c in str(sc)[:11]:
        n = n * 64 + ALFA.index(c)
    return n

def mid2sc(mid):
    mid = int(mid)
    s = ""
    while mid > 0:
        s = ALFA[mid % 64] + s
        mid //= 64
    return s

def sc_para_data(sc):
    ms = (sc2mid(sc) >> 23) + EPOCH_MS
    return datetime.datetime.utcfromtimestamp(ms / 1000.0 + OFFSET_S)

def mes_de(sc):
    return sc_para_data(sc).strftime("%Y-%m")

def _mes_rotulo(txt):
    txt = str(txt or "").strip()[:7]
    if len(txt) == 7 and txt[4] == "-" and txt[:4].isdigit() and txt[5:].isdigit():
        return txt
    return None

def _dia_rotulo(txt):
    txt = str(txt or "").strip()[:10]
    if (len(txt) == 10 and txt[4] == "-" and txt[7] == "-"
            and txt[:4].isdigit() and txt[5:7].isdigit() and txt[8:].isdigit()):
        return txt
    return None

# ------------------------------------------------------------------- leitores
def fonte_grid(caminho):
    doc = json.load(open(caminho, encoding="utf-8"))
    itens = doc["itens"] if isinstance(doc, dict) else doc
    out = {}
    for x in itens:
        sc = str(x.get("sc") or x.get("shortcode"))[:11]
        out[sc] = {"rotulo": str(x.get("publicado_brt") or "")[:16], "dono": x.get("dono_grid"), "tipo": x.get("tipo_grid")}
    return out

def fonte_jsonl(caminho):
    out = {}
    for linha in open(caminho, encoding="utf-8"):
        if linha.strip():
            x = json.loads(linha)
            out[str(x.get("shortcode"))[:11]] = {"rotulo": str(x.get("publicado_brt") or "")[:16], "dono": x.get("vinculo"), "tipo": x.get("tipo_midia")}
    return out

def fonte_brutos(pasta):
    out = {}
    for nome in os.listdir(pasta):
        if not nome.endswith(".json"):
            continue
        try:
            x = json.load(open(os.path.join(pasta, nome), encoding="utf-8"))
        except Exception:
            continue
        sc = str(x.get("shortcode") or nome[:-5])[:11]
        out[sc] = {"rotulo": "%s %s" % (x.get("data_publicacao") or "", x.get("hora_publicacao") or ""), "dono": x.get("autor"), "tipo": x.get("tipo_midia")}
    return out

def fonte_csv(caminho):
    out = {}
    for r in csv.DictReader(open(caminho, encoding="utf-8"), delimiter=";"):
        if not r.get("shortcode"):
            continue
        out[str(r["shortcode"])[:11]] = {"rotulo": "%s %s" % (r.get("data_publicacao") or "", r.get("hora_publicacao") or ""), "dono": r.get("vinculo"), "tipo": r.get("tipo_midia")}
    return out

def fonte_xlsx(caminho):
    import openpyxl
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = None
    for nome in ("Posts", "posts"):
        if nome in wb.sheetnames:
            ws = wb[nome]
            break
    ws = ws or wb.worksheets[0]
    cols = [c.value for c in ws[1]]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(cols, row))
        sc = r.get("shortcode")
        if not sc:
            continue
        out[str(sc)[:11]] = {"rotulo": str(r.get("data_publicacao") or ""), "dono": r.get("vinculo"), "tipo": r.get("tipo_midia")}
    return out

# Posts que NAO tem legenda por natureza: nao sao falha de captura. O texto do proprio
# post esta vazio, e o bruto do painel concorda. Ver
# auditorias/CONFERENCIA-API-LEGENDA-2026-08.md (post DcCPbyLGzGd, carrossel).
SEM_LEGENDA_POR_NATUREZA = {"DcCPbyLGzGd"}

def fonte_legendas(caminho):
    """Base de legendas: o mesmo conjunto de posts, com o texto do post.

    Coluna `legenda` obrigatoria daqui para frente (regra do LEIA-ME). Vazio = reprova.
    """
    out = {}
    for r in csv.DictReader(open(caminho, encoding="utf-8"), delimiter=";"):
        if not r.get("shortcode"):
            continue
        out[str(r["shortcode"])[:11]] = {
            "rotulo": "%s %s" % (r.get("data_publicacao") or "", r.get("hora_publicacao") or ""),
            "dono": r.get("vinculo"),
            "tipo": r.get("tipo_midia"),
            "legenda": (r.get("legenda") or "").strip(),
        }
    return out

def registro_fontes():
    """(leitor, caminho, descricao, classe, periodos_em_que_precisa_fechar).

    Classes: referencia (a varredura do feed), base (coluna publicada, nao pode
    ter rotulo errado), captura, export, varredura_antiga.
    """
    c = os.path.join(COFRE, "Hermes", "instagram-insights", "captura")
    r = os.path.join(COFRE, "000-Arquivos-originais", "Dados para mega-relatório pós Sudoexpo", "Redes sociais")
    b = os.path.join(COFRE, "85-Bases-e-Consultas")
    return {
        "grid": (fonte_grid, os.path.join(c, "setembro-2026", "_grid_inventario2.json"), "varredura do feed (02/10/2026)", "referencia", set()),
        "base_set": (fonte_csv, os.path.join(b, "Instagram-Publicacoes-2026-09.csv"), "base canonica de setembro", "base", {"2026-09"}),
        "base_ago26": (fonte_csv, os.path.join(b, "Instagram-Publicacoes-2026-08.csv"), "base canonica de agosto (62, completa)", "base", {"2026-08"}),
        "legendas_set": (fonte_legendas, os.path.join(b, "Instagram-Legendas-2026-09.csv"), "base de legendas de setembro (texto de cada post)", "legenda", {"2026-09"}),
        "legendas_ago": (fonte_legendas, os.path.join(b, "Instagram-Legendas-2026-08.csv"), "base de legendas de agosto (62, completa)", "legenda", {"2026-08"}),
        "base_08_09": (fonte_csv, os.path.join(b, "Instagram-Publicacoes-2026-08-09.csv"), "base antiga (agosto incompleto + setembro parcial) — superada por base_ago26 e base_set", "legado", set()),
        "captura_set": (fonte_jsonl, os.path.join(c, "setembro-2026", "_setembro_2026_116.jsonl"), "captura post-a-post de setembro", "captura", set()),
        "brutos_ago": (fonte_brutos, os.path.join(c, "agosto-2026"), "brutos post-a-post de agosto", "captura", set()),
        "export_set": (fonte_xlsx, os.path.join(r, "instagram_acirvoficial_insights_setembro_2026.xlsx"), "export oficial de setembro (so 01-14/09)", "export", set()),
        "export_ago": (fonte_xlsx, os.path.join(r, "instagram_acirvoficial_insights_agosto_2026_MASTER(1).xlsx"), "export oficial de agosto (MASTER)", "export", set()),
        "grid_v1": (fonte_grid, os.path.join(c, "setembro-2026", "_grid_inventario.json"), "varredura anterior (parcial)", "varredura_antiga", set()),
        "grid_v2": (fonte_grid, os.path.join(c, "setembro-2026", "_grid_inventario_v2.json"), "varredura anterior (parcial)", "varredura_antiga", set()),
    }

# ----------------------------------------------------------------- auditoria
def auditar(periodo, fontes=None, referencia="grid"):
    fontes = registro_fontes() if fontes is None else fontes
    dados = {}
    for nome, esp in fontes.items():
        leitor, caminho, desc, classe = esp[0], esp[1], esp[2], esp[3]
        exige = esp[4] if len(esp) > 4 else set()
        if caminho is not None and not os.path.exists(caminho):
            dados[nome] = {"desc": desc, "classe": classe, "exige": exige, "erro": "nao encontrado"}
            continue
        dados[nome] = {"desc": desc, "classe": classe, "exige": exige, "itens": leitor(caminho)}

    ref = dados[referencia]["itens"]
    ref_periodo = {sc for sc in ref if mes_de(sc) == periodo}
    res = {"periodo": periodo, "referencia": referencia, "total_referencia": len(ref_periodo),
           "fontes": {}, "divergencias": [], "avisos": [], "informacoes": []}

    for nome, d in dados.items():
        if "itens" not in d:
            res["fontes"][nome] = {"desc": d["desc"], "classe": d["classe"], "erro": d["erro"]}
            continue
        itens = d["itens"]
        no_periodo = {sc for sc in itens if mes_de(sc) == periodo}
        fora = sorted(no_periodo - ref_periodo)
        faltando = sorted(ref_periodo - no_periodo)
        rotulos = sorted(sc for sc in itens
                         if _mes_rotulo(d["itens"][sc].get("rotulo")) and _mes_rotulo(d["itens"][sc]["rotulo"]) != mes_de(sc)
                         and (mes_de(sc) == periodo or _mes_rotulo(d["itens"][sc]["rotulo"]) == periodo))
        dias = sorted(sc for sc in itens
                      if _dia_rotulo(d["itens"][sc].get("rotulo"))
                      and _dia_rotulo(d["itens"][sc]["rotulo"]) != sc_para_data(sc).strftime("%Y-%m-%d")
                      and _mes_rotulo(d["itens"][sc]["rotulo"]) == mes_de(sc) == periodo)
        sem_legenda = sorted(sc for sc in no_periodo
                             if not str(itens[sc].get("legenda") or "").strip()
                             and sc not in SEM_LEGENDA_POR_NATUREZA) \
            if d["classe"] == "legenda" else []
        res["fontes"][nome] = {"desc": d["desc"], "classe": d["classe"], "total": len(itens), "no_periodo": len(no_periodo),
                               "fora_da_referencia": fora, "faltando_no_periodo": faltando, "rotulo_errado": rotulos, "dia_errado": dias,
                               "sem_legenda": sem_legenda}
        if fora:
            res["divergencias"].append("%s traz %d post(s) que a varredura do feed nao tem: %s" % (nome, len(fora), ",".join(fora)))
        if rotulos:
            msg = "%s tem %d data(s) publicada(s) errada(s) para o periodo (a data do shortcode manda): %s" % (nome, len(rotulos), ",".join(rotulos))
            (res["divergencias"] if d["classe"] == "base" else res["avisos"]).append(msg)
        if dias:
            msg = "%s tem %d data(s) com mes certo e dia errado (o dia do shortcode manda): %s" % (nome, len(dias), ",".join(dias))
            (res["divergencias"] if d["classe"] == "base" else res["avisos"]).append(msg)
        if faltando and not fora and (no_periodo or periodo in d["exige"]):
            msg = "%s cobre %d de %d posts do periodo (faltam %d)" % (nome, len(no_periodo), len(ref_periodo), len(faltando))
            (res["divergencias"] if periodo in d["exige"] else res["informacoes"]).append(msg)
        if sem_legenda:
            msg = "%s tem %d post(s) de %s sem legenda: %s" % (nome, len(sem_legenda), periodo, ",".join(sem_legenda))
            (res["divergencias"] if d["classe"] == "legenda" else res["avisos"]).append(msg)
        if d["classe"] == "base" and periodo in d["exige"] and not no_periodo:
            res["divergencias"].append("%s nao tem nenhum post de %s" % (nome, periodo))

    res["veredito"] = "COMPLETO" if not res["divergencias"] else "REPROVADO"
    return res

def relatorio_md(res):
    L = ["# Auditoria de publicações — %s" % MESES.get(res["periodo"], res["periodo"]), "",
         "**Veredito: %s**" % res["veredito"], "",
         "Contagem do período = **%d posts**, definida pelo conjunto de shortcodes da varredura do feed (referência)." % res["total_referencia"], "",
         "| Fonte | O que é | Classe | Itens | No período | Fora da referência | Falta p/ fechar | Data publicada errada | Sem legenda |",
         "|---|---|---|---|---|---|---|---|---|"]
    for nome, f in res["fontes"].items():
        if "erro" in f:
            L.append("| `%s` | %s | %s | — | — | — | — | — | — |" % (nome, f["desc"], f["classe"]))
            continue
        L.append("| `%s` | %s | %s | %d | %d | %d | %d | %d | %d |" % (
            nome, f["desc"], f["classe"], f["total"], f["no_periodo"], len(f["fora_da_referencia"]), len(f["faltando_no_periodo"]), len(f["rotulo_errado"]) + len(f["dia_errado"]),
            len(f.get("sem_legenda") or [])))
    L.append("")
    for titulo, chave in (("Divergências (reprovam)", "divergencias"), ("Avisos", "avisos"), ("Informações", "informacoes")):
        if res[chave]:
            L.append("## %s" % titulo)
            for x in res[chave]:
                L.append("- %s" % x)
            L.append("")
    L.append("_Gerado por `Hermes/instagram-insights/ferramentas/publicacoes_feed.py`. "
             "Contagem = conjunto de shortcodes; data = decodificada do shortcode (BRT), nunca o rótulo da fonte._")
    return "\n".join(L) + "\n"

# ----------------------------------------------------------------- autoteste
def autoteste():
    """Bateria negativa com dados de mentira: prova que a auditoria reprova quando deve."""
    import tempfile
    tmp = tempfile.mkdtemp(prefix="audit_")
    try:
        g = fonte_grid(registro_fontes()["grid"][1])
    except Exception as e:
        return 0, ["nao consegui ler o grid real para montar a fixture: %s" % e]
    set_s = sorted(sc for sc in g if mes_de(sc) == "2026-09")
    S = set_s[:3]                       # 3 posts reais de setembro
    FANTASMA = set_s[3]                 # decodifica para setembro e NAO esta na fixture
    grid = os.path.join(tmp, "grid.json")
    json.dump({"itens": [{"sc": s, "publicado_brt": sc_para_data(s).strftime("%Y-%m-%d %H:%M:%S"), "dono_grid": "acirvoficial", "tipo_grid": "p"} for s in S]},
              open(grid, "w", encoding="utf-8"))

    def csv_de(nome, linhas, rotulos_errados=(), dias_errados=()):
        p = os.path.join(tmp, nome)
        with open(p, "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["shortcode", "data_publicacao", "hora_publicacao", "tipo_midia", "vinculo"])
            for s in linhas:
                d = sc_para_data(s)
                if s in rotulos_errados:
                    d = d - datetime.timedelta(days=40)
                if s in dias_errados:
                    mais = d + datetime.timedelta(days=1)
                    d = mais if mais.month == d.month else d - datetime.timedelta(days=1)
                w.writerow([s, d.strftime("%Y-%m-%d"), d.strftime("%H:%M:%S"), "p", "proprio"])
        return p

    base_ok = csv_de("base_ok.csv", S)
    base_falta = csv_de("base_falta.csv", S[:2])
    base_extra = csv_de("base_extra.csv", S + [FANTASMA])
    base_rotulo = csv_de("base_rotulo.csv", S, rotulos_errados={S[1]})
    base_dia = csv_de("base_dia.csv", S, dias_errados={S[1]})

    def montar(base, **extras):
        f = {"grid": (fonte_grid, grid, "grid de teste", "referencia", set()),
             "base": (fonte_csv, base, "base de teste", "base", {"2026-09"})}
        f.update(extras)
        return f

    export_ok = csv_de("export_ok.csv", S)
    export_rotulo = csv_de("export_rotulo.csv", S, rotulos_errados={S[2]})

    CASOS = [
        ("dados coerentes", montar(base_ok), "COMPLETO", None),
        ("base sem 1 post", montar(base_falta), "REPROVADO", "faltando"),
        ("base com post fantasma", montar(base_extra), "REPROVADO", "fantasma"),
        ("base com dia errado (mes certo)", montar(base_dia), "REPROVADO", "dia"),
        ("base com data publicada errada", montar(base_rotulo), "REPROVADO", "rotulo"),
        ("export com data errada (nao e base)", montar(base_ok, exp=(fonte_csv, export_rotulo, "export de teste", "export", set())), "COMPLETO", "rotulo"),
        ("export completo e coerente", montar(base_ok, exp=(fonte_csv, export_ok, "export de teste", "export", set())), "COMPLETO", None),
        ("varredura antiga com post que a nova nao tem", montar(base_ok, ant=(lambda p: dict(fonte_csv(base_ok), **{FANTASMA: {"rotulo": "2026-09-10", "dono": "x", "tipo": "p"}}), None, "antiga de teste", "varredura_antiga", set())), "REPROVADO", "fora"),
    ]

    ok, falhas = 0, []
    for nome, fonte, veredito, marca in CASOS:
        r = auditar("2026-09", fontes=fonte)
        certo = r["veredito"] == veredito
        if certo and marca == "faltando":
            certo = any("faltam" in x or "cobre" in x for x in r["divergencias"])
        if certo and marca == "dia":
            certo = any("dia errado" in x for x in r["divergencias"])
        if certo and marca == "fantasma":
            certo = any("nao tem" in x for x in r["divergencias"])
        if certo and marca == "rotulo":
            alvo = r["divergencias"] if veredito == "REPROVADO" else r["avisos"]
            certo = any("data" in x for x in alvo)
        if certo and marca == "fora":
            certo = any("nao tem" in x for x in r["divergencias"])
        if certo:
            ok += 1
        else:
            falhas.append("%s: veredito=%s (esperado %s)" % (nome, r["veredito"], veredito))
    return ok, falhas

# ------------------------------------------------------------------------ cli
def main(argv):
    global COFRE
    if "--cofre" in argv:
        COFRE = os.path.abspath(argv[argv.index("--cofre") + 1])
    if len(argv) > 1 and argv[1] == "autoteste":
        ok, falhas = autoteste()
        print("autoteste: %d/%d casos comportaram como esperado" % (ok, ok + len(falhas)))
        for f in falhas:
            print("   FALHA:", f)
        return 0 if not falhas else 1
    periodos = [argv[i + 1] for i, a in enumerate(argv) if a == "--periodo"] or ["2026-09", "2026-08"]
    destino = os.path.join(COFRE, "Hermes", "instagram-insights", "auditorias")
    os.makedirs(destino, exist_ok=True)
    pior = 0
    for p in periodos:
        res = auditar(p)
        json.dump(res, open(os.path.join(destino, "AUDITORIA-%s.json" % p), "w", encoding="utf-8"), ensure_ascii=False, indent=2)
        open(os.path.join(destino, "AUDITORIA-%s.md" % p), "w", encoding="utf-8").write(relatorio_md(res))
        print("== %s: %s | referencia = %d posts" % (p, res["veredito"], res["total_referencia"]))
        for n, f in res["fontes"].items():
            if "erro" in f:
                print("   %-12s -- %s" % (n, f["erro"]))
            else:
                print("   %-12s itens=%-4d no periodo=%-4d fora=%d falta=%d rotulo_errado=%d dia_errado=%d" % (
                    n, f["total"], f["no_periodo"], len(f["fora_da_referencia"]), len(f["faltando_no_periodo"]), len(f["rotulo_errado"]), len(f["dia_errado"])))
        for d in res["divergencias"]:
            print("   ! %s" % d)
        for a in res["avisos"]:
            print("   ~ %s" % a)
        for i in res["informacoes"]:
            print("   . %s" % i)
        pior = 1 if res["veredito"] != "COMPLETO" else pior
    print("relatorios: Hermes/instagram-insights/auditorias/AUDITORIA-<periodo>.md")
    return pior

if __name__ == "__main__":
    sys.exit(main(sys.argv))
